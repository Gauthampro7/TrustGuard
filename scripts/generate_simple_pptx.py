import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

# Colors (Light Theme)
COLOR_BG = RGBColor(248, 250, 252)         # #F8FAFC
COLOR_CARD_BG = RGBColor(255, 255, 255)    # #FFFFFF
COLOR_CARD_BORDER = RGBColor(226, 232, 240)# #E2E8F0
COLOR_TEXT_MAIN = RGBColor(15, 23, 42)     # #0F172A
COLOR_TEXT_MUTED = RGBColor(71, 85, 105)   # #475569
COLOR_TEXT_LIGHT = RGBColor(100, 116, 139) # #64748B
COLOR_ACCENT = RGBColor(37, 99, 235)       # #2563EB
COLOR_ACCENT_SOFT = RGBColor(239, 246, 255)# #EFF6FF
COLOR_ACCENT_BORDER = RGBColor(191, 219, 254) # #BFDBFE
COLOR_RED = RGBColor(239, 68, 68)          # #EF4444
COLOR_RED_BG = RGBColor(254, 242, 242)     # #FEF2F2
COLOR_GREEN = RGBColor(16, 185, 129)       # #10B981
COLOR_GREEN_BG = RGBColor(240, 253, 244)   # #F0FDF4
COLOR_AMBER = RGBColor(245, 158, 11)       # #F59E0B
COLOR_AMBER_BG = RGBColor(255, 251, 235)   # #FFFBEB

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

def set_slide_background(slide):
    bg_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = COLOR_BG
    bg_shape.line.fill.background()
    return bg_shape

def add_header(slide, eyebrow, title, subtitle=None):
    tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.35))
    tf = tx_box.text_frame
    tf.word_wrap = True
    tf.margin_top = tf.margin_bottom = tf.margin_left = tf.margin_right = 0
    p = tf.paragraphs[0]
    p.text = eyebrow.upper()
    p.font.size = Pt(10.5)
    p.font.bold = True
    p.font.color.rgb = COLOR_ACCENT

    tx_box2 = slide.shapes.add_textbox(Inches(0.8), Inches(0.72), Inches(11.7), Inches(0.55))
    tf2 = tx_box2.text_frame
    tf2.word_wrap = True
    tf2.margin_top = tf2.margin_bottom = tf2.margin_left = tf2.margin_right = 0
    p2 = tf2.paragraphs[0]
    p2.text = title
    p2.font.size = Pt(23)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_TEXT_MAIN

    if subtitle:
        tx_box3 = slide.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(11.7), Inches(0.35))
        tf3 = tx_box3.text_frame
        tf3.word_wrap = True
        tf3.margin_top = tf3.margin_bottom = tf3.margin_left = tf3.margin_right = 0
        p3 = tf3.paragraphs[0]
        p3.text = subtitle
        p3.font.size = Pt(12)
        p3.font.color.rgb = COLOR_TEXT_MUTED

def add_footer(slide, current_idx, total_count=11):
    tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.0), Inches(11.7), Inches(0.35))
    tf = tx_box.text_frame
    tf.margin_top = tf.margin_bottom = tf.margin_left = tf.margin_right = 0
    p = tf.paragraphs[0]
    p.text = f"TrustGuard · PS-02: AI for Digital Trust                                                                                                                      Slide {current_idx} of {total_count}"
    p.font.size = Pt(9.5)
    p.font.color.rgb = COLOR_TEXT_LIGHT

def add_card(slide, left, top, width, height, title, items, border_color=COLOR_CARD_BORDER, bg_color=COLOR_CARD_BG, top_bar=None):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    card.line.color.rgb = border_color
    card.line.width = Pt(1)

    if top_bar:
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(0.08))
        bar.fill.solid()
        bar.fill.fore_color.rgb = top_bar
        bar.line.fill.background()

    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.24)
    tf.margin_right = Inches(0.2)
    tf.margin_top = Inches(0.22)
    tf.margin_bottom = Inches(0.18)

    p0 = tf.paragraphs[0]
    p0.text = title
    p0.font.bold = True
    p0.font.size = Pt(13)
    p0.font.color.rgb = COLOR_TEXT_MAIN
    p0.space_after = Pt(8)

    for item in items:
        p = tf.add_paragraph()
        p.text = "• " + item if not item.startswith("  ") else item
        p.font.size = Pt(10.5)
        p.font.color.rgb = COLOR_TEXT_MUTED
        p.space_after = Pt(5)

def add_banner(slide, top, text):
    banner = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(top), Inches(11.7), Inches(0.75))
    banner.fill.solid()
    banner.fill.fore_color.rgb = COLOR_ACCENT_SOFT
    banner.line.color.rgb = COLOR_ACCENT_BORDER
    tf = banner.text_frame
    tf.margin_left = Inches(0.22)
    tf.margin_right = Inches(0.2)
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = COLOR_ACCENT

# ==================== SLIDE 1 ====================
s1 = prs.slides.add_slide(blank_layout)
set_slide_background(s1)

badge = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.5), Inches(1.8), Inches(4.3), Inches(0.42))
badge.fill.solid()
badge.fill.fore_color.rgb = COLOR_ACCENT_SOFT
badge.line.color.rgb = COLOR_ACCENT_BORDER
tf_b = badge.text_frame
p_b = tf_b.paragraphs[0]
p_b.text = "PROBLEM STATEMENT 02 · DIGITAL TRUST"
p_b.font.size = Pt(10.5)
p_b.font.bold = True
p_b.font.color.rgb = COLOR_ACCENT
p_b.alignment = PP_ALIGN.CENTER

t_box = s1.shapes.add_textbox(Inches(1.5), Inches(2.4), Inches(10.3), Inches(1.2))
tf_t = t_box.text_frame
p_t = tf_t.paragraphs[0]
p_t.text = "TrustGuard"
p_t.font.size = Pt(48)
p_t.font.bold = True
p_t.font.color.rgb = COLOR_TEXT_MAIN
p_t.alignment = PP_ALIGN.CENTER

st_box = s1.shapes.add_textbox(Inches(1.5), Inches(3.6), Inches(10.3), Inches(0.6))
tf_st = st_box.text_frame
p_st = tf_st.paragraphs[0]
p_st.text = "The All-in-One Digital Identity & Impersonation Defense Platform"
p_st.font.size = Pt(20)
p_st.font.bold = True
p_st.font.color.rgb = COLOR_ACCENT
p_st.alignment = PP_ALIGN.CENTER

d_box = s1.shapes.add_textbox(Inches(2.2), Inches(4.3), Inches(8.9), Inches(1.2))
tf_d = d_box.text_frame
tf_d.word_wrap = True
p_d = tf_d.paragraphs[0]
p_d.text = "A simple, unified web application and browser companion that helps you protect your online presence, hunt down fake accounts, detect unusual activity deviations, and evaluate profile trust in real-time."
p_d.font.size = Pt(13)
p_d.font.color.rgb = COLOR_TEXT_MUTED
p_d.alignment = PP_ALIGN.CENTER

f_box = s1.shapes.add_textbox(Inches(1.5), Inches(5.8), Inches(10.3), Inches(0.4))
tf_f = f_box.text_frame
p_f = tf_f.paragraphs[0]
p_f.text = "👤 Identity Baseline   •   🛡️ Fake Account Hunter   •   🧩 Browser Extension   •   🔍 Multimodal Inspector"
p_f.font.size = Pt(11.5)
p_f.font.color.rgb = COLOR_TEXT_LIGHT
p_f.alignment = PP_ALIGN.CENTER

# ==================== SLIDE 2 ====================
s2 = prs.slides.add_slide(blank_layout)
set_slide_background(s2)
add_header(s2, "01 · Why We Need This", "The Everyday Threat in 2026", "Scams today are no longer just bad Photoshop. They are coordinated, realistic identity attacks.")
add_card(s2, 0.8, 1.8, 3.7, 4.0, "🎭 Cloned Accounts & Lookalikes", [
    "Scammers copy your profile photo, bio, and name to create lookalike handles.",
    "Add sneaky suffixes like '_official' or '_support'.",
    "Substitute visually identical foreign letters (homoglyphs) to fool your colleagues and followers into sending money."
], border_color=COLOR_RED, bg_color=COLOR_RED_BG, top_bar=COLOR_RED)

add_card(s2, 4.8, 1.8, 3.7, 4.0, "🎙️ Voice & Video Impersonation", [
    "Just 10 seconds of clear speech from a video is enough for modern AI tools to clone anyone's voice.",
    "Scammers pair cloned voices with urgent WhatsApp calls or fake emergency family requests.",
    "Causes massive financial fraud and executive impersonation."
], border_color=COLOR_AMBER, bg_color=COLOR_AMBER_BG, top_bar=COLOR_AMBER)

add_card(s2, 8.8, 1.8, 3.7, 4.0, "🖼️ Stolen Content & Likeness", [
    "Creative professionals, designers, and influencers find their work stolen and reposted across the web.",
    "Watermarks are easily cropped out or filtered.",
    "Deepfakes and unauthorized AI models use artists' likeness without permission or licensing."
], border_color=COLOR_ACCENT, bg_color=COLOR_CARD_BG, top_bar=COLOR_ACCENT)

add_banner(s2, 6.0, "💡 Grammarly & QuillBot vs. TrustGuard: Tools like Grammarly or QuillBot check if a sentence has correct commas or if an essay was AI-generated. If an attacker hacks an account and writes a polite wire-transfer scam, Grammarly gives it a 100% green checkmark! Grammarly checks grammar; TrustGuard stops identity theft and fraud.")
add_footer(s2, 2)

# ==================== SLIDE 3 ====================
s3 = prs.slides.add_slide(blank_layout)
set_slide_background(s3)
add_header(s3, "02 · The Solution", "How TrustGuard Works in Practice", "A straightforward 4-step pipeline that turns scattered signals into clear, actionable trust.")

steps = [
    ("1", "Connect Profiles", "Add your official social accounts & samples"),
    ("2", "Build Baseline", "System learns your writing style & face/voice"),
    ("3", "Monitor & Detect", "Scans for fake accounts & abnormal posts"),
    ("4", "Explainable Report", "Get clear red flags & safe verification steps")
]
for i, (num, stitle, sdesc) in enumerate(steps):
    x_pos = Inches(0.8 + i * 2.95)
    node = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x_pos, Inches(1.8), Inches(2.7), Inches(1.3))
    node.fill.solid()
    node.fill.fore_color.rgb = COLOR_CARD_BG
    node.line.color.rgb = COLOR_ACCENT if i == 2 else COLOR_CARD_BORDER
    tf = node.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    p1.text = f"Step {num}: {stitle}"
    p1.font.bold = True
    p1.font.size = Pt(11)
    p1.font.color.rgb = COLOR_TEXT_MAIN
    p1.alignment = PP_ALIGN.CENTER
    p2 = tf.add_paragraph()
    p2.text = sdesc
    p2.font.size = Pt(9.5)
    p2.font.color.rgb = COLOR_TEXT_MUTED
    p2.alignment = PP_ALIGN.CENTER

add_card(s3, 0.8, 3.4, 5.7, 3.3, "🖥️ Unified Web Dashboard", [
    "User logs in and connects their social media footprint in one unified dashboard.",
    "Sets up an automated security shield that watches for lookalikes and copycats across platforms.",
    "Shows clean security scores, clone alerts, and unauthorized content detections."
], border_color=COLOR_CARD_BORDER, bg_color=COLOR_CARD_BG, top_bar=COLOR_ACCENT)

add_card(s3, 6.8, 3.4, 5.7, 3.3, "🧩 In-Browser Companion & Inspector", [
    "Works everywhere you browse: actively checks profiles as you open them on Instagram, X, or LinkedIn.",
    "Calculates a continuous authenticity score directly in your feed.",
    "Lets you drop in suspicious media files for instant explainable breakdown."
], border_color=COLOR_CARD_BORDER, bg_color=COLOR_CARD_BG, top_bar=COLOR_GREEN)
add_footer(s3, 3)

# ==================== SLIDE 4 ====================
s4 = prs.slides.add_slide(blank_layout)
set_slide_background(s4)
add_header(s4, "03 · Practical Feature 1", "Your Personal Identity Baseline", "Building a secure reference summary so the system knows what is truly 'you'.")

add_card(s4, 0.8, 1.8, 5.7, 4.0, "📂 What the User Provides", [
    "Official Account Links: Connect verified LinkedIn, X/Twitter, YouTube, GitHub, or portfolio link.",
    "Face & Photo References: Upload 3–5 clear reference photos to register your authentic face appearance.",
    "Voice Sample (Optional): A quick 30-second audio clip reading a sentence to register your unique voice tone.",
    "Writing Style Samples: A sample of past genuine posts/articles to establish your normal topics and writing style."
], border_color=COLOR_ACCENT, bg_color=COLOR_CARD_BG, top_bar=COLOR_ACCENT)

add_card(s4, 6.8, 1.8, 5.7, 4.0, "🔒 Privacy First (Zero Data Selling)", [
    "No Raw Biometric Storage: Photos and audio clips are converted into mathematical fingerprint numbers. Raw images are never stored or sold.",
    "Zero Model Training: Your personal data is never used to train public generative AI models.",
    "Encrypted & Controlled by You: Data is encrypted locally and in transit.",
    "1-Click Erasure: Delete your profile and permanently wipe all baseline fingerprints anytime."
], border_color=COLOR_GREEN, bg_color=COLOR_GREEN_BG, top_bar=COLOR_GREEN)

add_banner(s4, 6.0, "💡 Why this matters: You cannot detect an 'impersonation' without knowing the authentic person. This baseline gives the system the ground truth it needs to protect you.")
add_footer(s4, 4)

# ==================== SLIDE 5 ====================
s5 = prs.slides.add_slide(blank_layout)
set_slide_background(s5)
add_header(s5, "04 · Practical Feature 2", "Fake Account & Clone Hunter", "Automatically finding people pretending to be you across social platforms.")

add_card(s5, 0.8, 1.8, 3.7, 4.0, "1. Multi-Tier Handle Scanner", [
    "Catches deceptive impersonation variations:",
    "  • Separator padding: e.g. adding extra underscores 'user______name' vs 'user_name'.",
    "  • Typosquat mutations: Levenshtein distance ≤ 3 edits.",
    "  • Lookalike letters (homoglyphs): Replacing Latin 'a' with Cyrillic 'а' (TR39 skeleton collision).",
    "  • Contextual mismatches: Unrelated handles marked neutral, avoiding false accusations."
], border_color=COLOR_AMBER, bg_color=COLOR_AMBER_BG, top_bar=COLOR_AMBER)

add_card(s5, 4.8, 1.8, 3.7, 4.0, "2. Profile Picture & Bio Match", [
    "Searches platforms for accounts that have:",
    "  • Copied your exact profile photo or cropped versions of it.",
    "  • Copied your bio text, company name, or professional title.",
    "  • Unlinked accounts attempting to steal your professional credibility."
], border_color=COLOR_ACCENT, bg_color=COLOR_CARD_BG, top_bar=COLOR_ACCENT)

add_card(s5, 8.8, 1.8, 3.7, 4.0, "3. Scam Activity Alert", [
    "Flags malicious behaviors:",
    "  • Recently created clone accounts messaging your real connections.",
    "  • Profiles promoting crypto giveaways, fake jobs, or unapproved paid services in your name."
], border_color=COLOR_RED, bg_color=COLOR_RED_BG, top_bar=COLOR_RED)

add_banner(s5, 6.0, "📋 Actionable Output: TrustGuard gives you a direct 1-click Takedown & Report Package with pre-filled evidence links to submit to Twitter, LinkedIn, or Meta support.")
add_footer(s5, 5)

# ==================== SLIDE 6 ====================
s6 = prs.slides.add_slide(blank_layout)
set_slide_background(s6)
add_header(s6, "05 · Practical Feature 3", "Activity & Post Anomaly Detection", "Detecting when an account is hacked or someone is speaking completely out-of-character.")

add_card(s6, 0.8, 1.8, 5.7, 4.0, "🚨 What an Anomaly Looks Like", [
    "Tone & Topic Shift: A software developer who normally posts about coding suddenly promotes an urgent cryptocurrency investment scheme.",
    "Coercion & Panic Style: Sudden urgent, demanding tone: 'URGENT: Emergency wire needed! Do not call me, phone line is broken!'",
    "Unusual Timing: Strange post bursts at 3:30 AM local time or 15 automated posts in 2 minutes.",
    "Life Summary Mismatch: Post claims to be stuck in a hospital abroad while official calendar confirms they are at their office desk."
], border_color=COLOR_AMBER, bg_color=COLOR_AMBER_BG, top_bar=COLOR_AMBER)

add_card(s6, 6.8, 1.8, 5.7, 4.0, "⚙️ How the System Detects It", [
    "Writing Habit Comparison: Checks sentence lengths, vocabulary variety, and punctuation habits against your baseline writeprint.",
    "AI Text Detection: Flags if the post reads like a generic template written by ChatGPT or automated bots.",
    "Urgency Scoring: Detects psychological pressure tactics commonly used in phishing, CEO fraud, and extortion.",
    "Automated Warning: Immediately alerts family, team members, or followers before funds are transferred."
], border_color=COLOR_ACCENT, bg_color=COLOR_CARD_BG, top_bar=COLOR_ACCENT)

add_banner(s6, 6.0, "🛡️ Real-World Use: Protects organizations and families from the 'CEO Fraud' or 'Emergency Family Scam' where attackers spoof a persona to steal urgent money.")
add_footer(s6, 6)

# ==================== SLIDE 7 ====================
s7 = prs.slides.add_slide(blank_layout)
set_slide_background(s7)
add_header(s7, "06 · Practical Feature 4", "Copyright & Likeness Shield", "Protecting creative professionals, designers, and influencers from unauthorized reuse.")

add_card(s7, 0.8, 1.8, 3.7, 4.0, "🎨 Artwork & Photo Matching", [
    "Artists upload original designs, photos, or digital art.",
    "System catches reposts across the web even if the thief:",
    "  • Cropped out the artist's watermark.",
    "  • Applied color filters or flipped the image.",
    "  • Turned it into an unauthorized print or NFT."
], border_color=COLOR_ACCENT, bg_color=COLOR_CARD_BG, top_bar=COLOR_ACCENT)

add_card(s7, 4.8, 1.8, 3.7, 4.0, "🎙️ Voice Likeness Defense", [
    "For podcasters, voice actors, and speakers:",
    "  • Alerts when your unique voice timbre is detected inside unauthorized commercial ads or AI voice apps.",
    "  • Confirms whether an audio clip is authentic or synthesized by generative tools."
], border_color=COLOR_ACCENT, bg_color=COLOR_CARD_BG, top_bar=COLOR_ACCENT)

add_card(s7, 8.8, 1.8, 3.7, 4.0, "📜 Ownership Provenance", [
    "Integrates Content Credentials (C2PA):",
    "  • Verifies cryptographic authorship certificates from cameras and Adobe tools.",
    "  • Confirms if media was taken with a real camera vs. created with Midjourney or DALL-E."
], border_color=COLOR_GREEN, bg_color=COLOR_GREEN_BG, top_bar=COLOR_GREEN)

add_banner(s7, 6.0, "💼 For Professionals: Creates a ready-to-file digital copyright dossier that creators can use to file DMCA takedown requests or claim licensing royalties.")
add_footer(s7, 7)

# ==================== SLIDE 8: EXTENSION (NEW) ====================
s8 = prs.slides.add_slide(blank_layout)
set_slide_background(s8)
add_header(s8, "07 · Practical Feature 5", "The Always-On Browser Extension", "Real-time profile genuineness as you browse Instagram, X, or LinkedIn.")

add_card(s8, 0.8, 1.8, 5.7, 4.0, "🌐 3-Platform DOM Adapters & Multi-Image Screening", [
    "Native Profile Inspection: Directly supports public profiles on Instagram, X/Twitter, and LinkedIn.",
    "Social Footprint Parsing: Reads posts count, followers, and following directly from the page layout.",
    "Multi-Image AI Screening: Samples avatar and up to 3 post thumbnails, screening 2D-FFT spectra for synthetic lattice artifacts.",
    "Memory Safety & Privacy: Ephemeral 64x64 downsamples strictly disposed in finally blocks; zero raw media or passwords ever stored.",
    "SPA Route Protection: Automatically resets reference handles and observation state during route changes."
], border_color=COLOR_ACCENT, bg_color=COLOR_CARD_BG, top_bar=COLOR_ACCENT)

add_card(s8, 6.8, 1.8, 5.7, 4.0, "📈 Continuous Calibrated Authenticity Index (0 - 100)", [
    "Parametric model replacing crude binary flags with real profile variance:",
    "  • Logarithmic Post Volume: Continuous scaling from 1 to 100+ posts.",
    "  • Network Depth: Logarithmic scaling from 10 to 5,000+ followers.",
    "  • Reciprocity Ratios: Rewards mutual networks; penalizes follow-churn bots (ratio > 40).",
    "  • Optical Texture: Measures avatar pixel standard deviation to distinguish real photos from blanks.",
    "  • Real Examples: Active creator (91/100), typical user (82/100), new burner (60/100), Cyrillic lookalike (5/100)."
], border_color=COLOR_GREEN, bg_color=COLOR_CARD_BG, top_bar=COLOR_GREEN)

add_banner(s8, 6.0, "⚡ Zero External API Costs: By inspecting public DOM elements on-demand, TrustGuard requires zero expensive platform APIs, operates 100% locally, and respects platform privacy boundaries.")
add_footer(s8, 8)

# ==================== SLIDE 9: MULTIMODAL INSPECTOR ====================
s9 = prs.slides.add_slide(blank_layout)
set_slide_background(s9)
add_header(s9, "08 · Practical Feature 6", "The Multimodal Inspector & ML-1", "Pretrained AI classifiers, signal-bounded uncertainty, and explainable dual ledgers.")

add_card(s9, 0.8, 1.8, 5.7, 2.5, "📥 Multimodal Forensic Suite", [
    "Nine CPU Extractors & Two Pretrained AI Models:",
    "  • RoBERTa AI-text classifier & Swin-v2 AI-image classifier (local, air-gapped).",
    "  • Voice vocoder artifact detection & cross-modal A/V sync tracking.",
    "  • 2D-FFT radial spectral rolloff & TR39 homoglyph Unicode hunter."
], border_color=COLOR_CARD_BORDER, bg_color=COLOR_CARD_BG, top_bar=COLOR_ACCENT)

add_card(s9, 6.8, 1.8, 5.7, 2.5, "🔍 Signal-Bounded Epistemic Uncertainty", [
    "Key Innovation: Prevents bundle-wide dilution:",
    "  • Each extractor signal is bounded by its own usability (min(score, 1.15 - U)).",
    "  • A short bio or missing modality no longer outvotes a clear forensic red flag.",
    "  • Evaluated uncertainty computed over active extractors only."
], border_color=COLOR_ACCENT, bg_color=COLOR_CARD_BG, top_bar=COLOR_ACCENT)

res_box = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.5), Inches(11.7), Inches(2.2))
res_box.fill.solid()
res_box.fill.fore_color.rgb = COLOR_CARD_BG
res_box.line.color.rgb = COLOR_CARD_BORDER
tf_res = res_box.text_frame
tf_res.margin_left = Inches(0.24)
tf_res.margin_top = Inches(0.18)
p_r1 = tf_res.paragraphs[0]
p_r1.text = "📊 Dual Evidence Ledger: Separating Red Flags from Uncertainty"
p_r1.font.bold = True
p_r1.font.size = Pt(12)
p_r1.font.color.rgb = COLOR_TEXT_MAIN

rc = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.1), Inches(5.1), Inches(5.3), Inches(1.3))
rc.fill.solid()
rc.fill.fore_color.rgb = COLOR_RED_BG
rc.line.color.rgb = COLOR_RED
tf_rc = rc.text_frame
tf_rc.margin_left = Inches(0.15)
p_rc = tf_rc.paragraphs[0]
p_rc.text = "🔴 Forensic Red Flags (Definite Anomalies)"
p_rc.font.bold = True
p_rc.font.size = Pt(11)
p_rc.font.color.rgb = COLOR_RED
p_rc2 = tf_rc.add_paragraph()
p_rc2.text = "• Audio vocoder synthesis signatures detected.\n• Lip-sync desynchronization lag > 180 ms.\n• Observed handle contains Cyrillic homoglyph (U+0430)."
p_rc2.font.size = Pt(9.5)
p_rc2.font.color.rgb = COLOR_TEXT_MUTED

gc = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.7), Inches(5.1), Inches(5.5), Inches(1.3))
gc.fill.solid()
gc.fill.fore_color.rgb = COLOR_ACCENT_SOFT
gc.line.color.rgb = COLOR_ACCENT_BORDER
tf_gc = gc.text_frame
tf_gc.margin_left = Inches(0.15)
p_gc = tf_gc.paragraphs[0]
p_gc.text = "🔍 Uncertainty & Context (Honest Limits)"
p_gc.font.bold = True
p_gc.font.size = Pt(11)
p_gc.font.color.rgb = COLOR_ACCENT
p_gc2 = tf_gc.add_paragraph()
p_gc2.text = "• Missing modalities reported as unavailable, never zero-risk.\n• Polished text shown as context, not false red flags.\n• Step-by-step verification playbook provided for analysts."
p_gc2.font.size = Pt(9.5)
p_gc2.font.color.rgb = COLOR_TEXT_MUTED
add_footer(s9, 9)

# ==================== SLIDE 10: NO SCORE IS PROOF ====================
s10 = prs.slides.add_slide(blank_layout)
set_slide_background(s10)
add_header(s10, "09 · Core Hackathon Ethos", "Why \"No Score is Proof\"", "An AI score is never absolute truth. TrustGuard gives you safe ways to verify independently.")

add_card(s10, 0.8, 1.8, 3.7, 4.0, "1. 5D Calibrated Trust Vector", [
    "Replaces misleading scalar ratings with five explicit dimensions:",
    "  • Media Anomaly (synthesis)",
    "  • Cross-Modal Discordance (lip-sync)",
    "  • Identity Mismatch (homoglyphs)",
    "  • Context Anomaly (urgency/style)",
    "  • Epistemic Uncertainty (data quality)"
], border_color=COLOR_ACCENT, bg_color=COLOR_CARD_BG, top_bar=COLOR_ACCENT)

add_card(s10, 4.8, 1.8, 3.7, 4.0, "2. Independent Verification", [
    "Instead of leaving users helpless, TrustGuard provides:",
    "  • Ground-truth company directory lookup.",
    "  • Steganographic canary tripwire checking.",
    "  • Adversarial dialectic: provides explicit arguments for AND against authenticity."
], border_color=COLOR_GREEN, bg_color=COLOR_GREEN_BG, top_bar=COLOR_GREEN)

add_card(s10, 8.8, 1.8, 3.7, 4.0, "3. Signed Audit Certificates", [
    "Cryptographically seals the final determination:",
    "  • Ed25519 digital signature with embedded QR verification routing.",
    "  • Verifiable offline or online with SHA-256 evidence hashes.",
    "  • Downloadable as official PDF for banking, IT, or legal compliance."
], border_color=COLOR_ACCENT, bg_color=COLOR_CARD_BG, top_bar=COLOR_ACCENT)

add_banner(s10, 6.0, "🎯 Judges' Favorite: The official hackathon problem explicitly states: 'Teams must demonstrate that the system does not treat an AI score as proof.' TrustGuard enforces this in code, UI meters, and audit trails.")
add_footer(s10, 10)

# ==================== SLIDE 11: DEMO SCOPE ====================
s11 = prs.slides.add_slide(blank_layout)
set_slide_background(s11)
add_header(s11, "10 · Verified Working Prototype", "Fully Implemented & Battle-Tested", "All components live, tested across 414 test gates, and demonstrable on stage.")

add_card(s11, 0.8, 1.8, 5.7, 4.0, "💻 Live Working Components (Delivered)", [
    "Interactive Web App Dashboard: 5D vector meters, dual evidence ledgers, and Reversible Adversarial Sandbox (AR-1).",
    "Live Browser Extension: Real-time public profile inspection on Instagram, X, and LinkedIn with continuous dynamic index.",
    "Multi-Modal Pipeline: 9 CPU extractors (<27ms median) + optional local ML classifiers (<250ms p95).",
    "414 Automated Checks Passing: 100% test pass rate across unit, integration, extension, and real headless Chromium tests."
], border_color=COLOR_ACCENT, bg_color=COLOR_CARD_BG, top_bar=COLOR_ACCENT)

add_card(s11, 6.8, 1.8, 5.7, 4.0, "🎬 4 Complete Production Scenarios", [
    "Scenario 1: CEO Wire Scam: Cloned voice + lip-sync mismatch + urgent text -> HIGH_IMPERSONATION_RISK.",
    "Scenario 2: Homoglyph Clone: Unicode TR39 spoofing + pHash match -> HIGH_IMPERSONATION_RISK.",
    "Scenario 3: Creator Copyright & Canary: Steganographic invisible zero-width marker triggers live tripwire alert.",
    "Scenario 4: Wi-Fi Compression Edge Case: Degraded channel noise cleanly clamps uncertainty to prevent false accusation."
], border_color=COLOR_GREEN, bg_color=COLOR_GREEN_BG, top_bar=COLOR_GREEN)

add_banner(s11, 6.0, "🚀 Conclusion: TrustGuard delivers digital trust without data hoarding, provides actionable proof instead of AI guesswork, and is 100% ready for real-world deployment.")
add_footer(s11, 11)

# Save
base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
pres_dir = os.path.join(base_dir, "docs", "presentations")
os.makedirs(pres_dir, exist_ok=True)
out_pptx = os.path.join(pres_dir, "presentation_simple.pptx")
prs.save(out_pptx)
print("Saved simple PPTX to:", out_pptx)
